import io
from pathlib import Path

def parse_pdf(file_bytes):
    import fitz
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    out = []
    for i, page in enumerate(doc):
        out.append({"page": i+1, "text": (page.get_text("text") or "").strip()})
    return out

def parse_docx(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text.strip() for c in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return [{"page": None, "text": "\n".join(parts)}]

def parse_file(filename, file_bytes):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_bytes)
    if ext == ".docx":
        return parse_docx(file_bytes)
    if ext in {".txt",".md"}:
        return [{"page": None, "text": file_bytes.decode("utf-8", errors="ignore")}]
    raise ValueError(f"지원하지 않는 파일 형식: {ext}")

def pages_to_text(pages):
    return "\n\n".join(
        f"[PAGE {p.get('page') if p.get('page') else '-'}]\n{p.get('text','')}"
        for p in pages
    )

def is_image_pdf(pages):
    if not pages:
        return False
    nontrivial = sum(1 for p in pages if len((p.get("text") or "").strip()) >= 50)
    return nontrivial < max(1, len(pages) // 3)

def chunks_from_pages(pages, max_chars=38000):
    chunks = []
    cur = []
    size = 0
    for p in pages:
        block = f"[PAGE {p.get('page') if p.get('page') else '-'}]\n{p.get('text','')}\n"
        if cur and size + len(block) > max_chars:
            chunks.append("\n".join(cur))
            cur, size = [], 0
        cur.append(block)
        size += len(block)
    if cur:
        chunks.append("\n".join(cur))
    return chunks
